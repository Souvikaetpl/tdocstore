import logging
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

import docx
import openpyxl
from pptx import Presentation
from pypdf import PdfReader

from . import config

log = logging.getLogger("crawler.extract")

_PREFERRED_EXTENSIONS = [".docx", ".pdf", ".pptx", ".xlsx", ".doc", ".ppt", ".xls", ".txt"]

# .doc/.ppt need a real file on disk and a LibreOffice subprocess, unlike
# every other format here which reads straight from an open zip member —
# handled separately in extract_text_from_zip rather than through
# _EXTRACTORS below.
_LIBREOFFICE_CONVERTIBLE = {".doc", ".ppt"}


@dataclass
class ExtractionResult:
    status: str  # 'success' | 'unsupported' | 'no_document' | 'error'
    text: Optional[str] = None
    source_filename: Optional[str] = None
    error: Optional[str] = None


def pick_main_entry(names: list[str], tdoc_id: str) -> Optional[str]:
    candidates = [n for n in names if Path(n).suffix.lower() in _PREFERRED_EXTENSIONS]
    if not candidates:
        return None

    prefix_matches = [n for n in candidates if Path(n).stem.lower().startswith(tdoc_id.lower())]
    pool = prefix_matches or candidates

    def rank(name: str) -> int:
        ext = Path(name).suffix.lower()
        return _PREFERRED_EXTENSIONS.index(ext) if ext in _PREFERRED_EXTENSIONS else len(_PREFERRED_EXTENSIONS)

    return sorted(pool, key=rank)[0]


def _dedupe_consecutive(items: list[str]) -> list[str]:
    """python-docx repeats a merged cell's text once per grid column it
    spans, so a title cell merged across 7 columns yields the same text
    7 times in row.cells. Collapse consecutive repeats back to one."""
    out: list[str] = []
    for item in items:
        if not out or out[-1] != item:
            out.append(item)
    return out


def _extract_docx(fileobj) -> str:
    document = docx.Document(fileobj)
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = _dedupe_consecutive([c.text.strip() for c in row.cells if c.text.strip()])
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_pdf(fileobj) -> str:
    reader = PdfReader(fileobj)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_pptx(fileobj) -> str:
    presentation = Presentation(fileobj)
    parts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
    return "\n".join(parts)


def _extract_xlsx(fileobj) -> str:
    wb = openpyxl.load_workbook(fileobj, data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(v) for v in row if v is not None]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def _extract_txt(fileobj) -> str:
    return fileobj.read().decode("utf-8", errors="replace")


_EXTRACTORS = {
    ".docx": _extract_docx,
    ".pdf": _extract_pdf,
    ".pptx": _extract_pptx,
    ".xlsx": _extract_xlsx,
    ".txt": _extract_txt,
}


def convert_doc_to_txt(path: Path) -> Optional[str]:
    """Legacy binary .doc -> plain text via LibreOffice headless conversion.
    python-docx is .docx-only, so this is the only way to read the older
    format. Also reused by meeting_info.py for Invitation documents."""
    if config.SOFFICE_PATH is None:
        return None
    with tempfile.TemporaryDirectory() as tmp_dir:
        result = subprocess.run(
            [config.SOFFICE_PATH, "--headless", "--convert-to", "txt", "--outdir", tmp_dir, str(path)],
            capture_output=True, text=True, timeout=config.RENDER_TIMEOUT_SECONDS,
        )
        out_path = Path(tmp_dir) / (path.stem + ".txt")
        if out_path.exists():
            return out_path.read_text(encoding="utf-8", errors="replace")
        log.warning("soffice txt conversion failed for %s: %s", path, result.stderr or result.stdout)
        return None


def _convert_ppt_to_text(path: Path) -> Optional[str]:
    """Legacy binary .ppt -> text via LibreOffice. Converts to .pptx first
    rather than trusting soffice's own txt export for presentations (which
    only pulls outline/notes text, not text sitting in arbitrary on-slide
    shapes) — then reuses the existing shape-by-shape _extract_pptx."""
    if config.SOFFICE_PATH is None:
        return None
    with tempfile.TemporaryDirectory() as tmp_dir:
        result = subprocess.run(
            [config.SOFFICE_PATH, "--headless", "--convert-to", "pptx", "--outdir", tmp_dir, str(path)],
            capture_output=True, text=True, timeout=config.RENDER_TIMEOUT_SECONDS,
        )
        out_path = Path(tmp_dir) / (path.stem + ".pptx")
        if not out_path.exists():
            log.warning("soffice pptx conversion failed for %s: %s", path, result.stderr or result.stdout)
            return None
        with open(out_path, "rb") as f:
            return _extract_pptx(f)


def extract_text_from_zip(zip_path: Path, tdoc_id: str) -> ExtractionResult:
    try:
        with zipfile.ZipFile(zip_path) as z:
            names = z.namelist()
            main_entry = pick_main_entry(names, tdoc_id)
            if main_entry is None:
                return ExtractionResult(status="no_document")

            ext = Path(main_entry).suffix.lower()

            if ext in _LIBREOFFICE_CONVERTIBLE:
                if config.SOFFICE_PATH is None:
                    return ExtractionResult(status="unsupported", source_filename=main_entry)
                with tempfile.TemporaryDirectory() as tmp_dir:
                    tmp_path = Path(tmp_dir) / Path(main_entry).name
                    with z.open(main_entry) as src, open(tmp_path, "wb") as dst:
                        shutil.copyfileobj(src, dst)
                    text = convert_doc_to_txt(tmp_path) if ext == ".doc" else _convert_ppt_to_text(tmp_path)
                if text is None:
                    return ExtractionResult(
                        status="error", source_filename=main_entry,
                        error="LibreOffice conversion produced no output",
                    )
                return ExtractionResult(status="success", text=text, source_filename=main_entry)

            extractor = _EXTRACTORS.get(ext)
            if extractor is None:
                return ExtractionResult(status="unsupported", source_filename=main_entry)

            with z.open(main_entry) as fileobj:
                text = extractor(fileobj)
            return ExtractionResult(status="success", text=text, source_filename=main_entry)
    except zipfile.BadZipFile as exc:
        return ExtractionResult(status="error", error=f"bad zip: {exc}")
    except Exception as exc:
        log.exception("Extraction failed for %s (%s)", zip_path, tdoc_id)
        return ExtractionResult(status="error", error=str(exc))
