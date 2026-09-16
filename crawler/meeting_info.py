"""Parse a meeting's Invitation document for its real date range and
location. Neither the TDoc list Excel nor any document's own cover page
carries this — it only exists in the (often multi-group, joint-venue)
Invitation document, as a structured table right at the top:

    TITLE          TYPE   DATES                  LOCATION   CTRY
    3GPPCT1#162    OR     24 - 28 August 2026     Prague     Czech Republic

One invitation commonly covers several co-located meetings (e.g. a joint
CT+SA week) at once; this extracts only the row matching the specific
meeting being ingested.
"""
import logging
import re
import shutil
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from .extract import _extract_pdf, convert_doc_to_txt

_INVITATION_EXTENSIONS = (".doc", ".docx", ".pdf")

log = logging.getLogger("crawler.meeting_info")

_MONTHS = {
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
    "july": 7, "august": 8, "september": 9, "october": 10, "november": 11, "december": 12,
    # Abbreviated form, observed in at least one invitation ("17 - 18 Jun 2024")
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "jun": 6, "jul": 7, "aug": 8,
    "sep": 9, "sept": 9, "oct": 10, "nov": 11, "dec": 12,
}

# A row splits into >=4 whitespace-run-separated fields once tabs (or
# multi-space alignment, depending on the export) are collapsed.
_SPLIT_RE = re.compile(r"\t+| {2,}")
# The day-range separator: hyphen or en-dash normally, but pypdf sometimes
# mis-decodes an invitation's own dash glyph from a broken/custom PDF font
# encoding into the literal Unicode replacement character (observed: SA4#136's
# "May 11 <?> 15, 2026") — tolerated here rather than losing the date entirely.
_DASH = "[-–�]"
# ",?" before the year: some invitations write "10 September, 2024" (comma
# before the year), others "28 August 2026" (no comma) — both observed.
_SAME_MONTH_RE = re.compile(rf"(\d{{1,2}})\s*{_DASH}\s*(\d{{1,2}})\s+([A-Za-z]+),?\s+(\d{{4}})")
_CROSS_MONTH_RE = re.compile(rf"(\d{{1,2}})\s+([A-Za-z]+)\s*{_DASH}\s*(\d{{1,2}})\s+([A-Za-z]+),?\s+(\d{{4}})")
# Month-first variant, observed in a RAN1 "-bis" invitation: "Oct 12 - 16, 2026".
_MONTH_FIRST_RE = re.compile(rf"([A-Za-z]+)\s+(\d{{1,2}})\s*{_DASH}\s*(\d{{1,2}}),?\s+(\d{{4}})")
# Non-place cells sometimes baked into the table in place of a real
# location column (observed: a "Register Here" hyperlink label). Word-
# bounded and specific on purpose — a generic "link" substring check
# would misfire on a real city name like "Linkoping".
_JUNK_CELL_RE = re.compile(r"\bregister\b|\bclick here\b|\bN/?A\b|\bTBD\b", re.IGNORECASE)


@dataclass
class MeetingInfo:
    start_date: Optional[str] = None
    end_date: Optional[str] = None
    location: Optional[str] = None
    country: Optional[str] = None


def _parse_date_range(text: str):
    m = _CROSS_MONTH_RE.search(text)
    if m:
        day1, mon1, day2, mon2, year = m.groups()
        mon1_num, mon2_num = _MONTHS.get(mon1.lower()), _MONTHS.get(mon2.lower())
        if mon1_num and mon2_num:
            return (f"{year}-{mon1_num:02d}-{int(day1):02d}", f"{year}-{mon2_num:02d}-{int(day2):02d}")
    m = _SAME_MONTH_RE.search(text)
    if m:
        day1, day2, mon, year = m.groups()
        mon_num = _MONTHS.get(mon.lower())
        if mon_num:
            return (f"{year}-{mon_num:02d}-{int(day1):02d}", f"{year}-{mon_num:02d}-{int(day2):02d}")
    m = _MONTH_FIRST_RE.search(text)
    if m:
        mon, day1, day2, year = m.groups()
        mon_num = _MONTHS.get(mon.lower())
        if mon_num:
            return (f"{year}-{mon_num:02d}-{int(day1):02d}", f"{year}-{mon_num:02d}-{int(day2):02d}")
    return (None, None)


def _normalize_title(s: str) -> str:
    """Strips whitespace AND hyphens, uppercases — makes "3GPP CT#105"
    (space), "RAN1#126-bis" (hyphen) and "RAN1#126bis" (no separator) all
    compare equal, without needing to predict which spelling a given
    invitation happens to use."""
    return re.sub(r"[\s\-]+", "", s).upper()


# A joint TSG-plenary week invitation (ATIS-hosted, observed for CT/RAN/SA
# plenary meetings in North America) lists each TSG on its own line as
# "CT#114 December 7-8, 2026" — a bare title then a date range, single-
# space separated (no table columns at all), followed a couple of lines
# later by a standalone "City, ST, Country" location line shared by all
# the TSGs listed above it.
_TITLE_DATE_LINE_RE = re.compile(r"^([A-Za-z0-9]+#\d+[A-Za-z\-]*)\s+(.+?)\s*$")

# A joint WG-week invitation (ATIS-hosted "3GPP WORKING GROUP MEETINGS
# INVITATION" letter) has no per-group breakdown at all — one shared
# date/location sentence ("...which will be held November 16-20, 2026, in
# Calgary, Alberta, Canada.") followed by a plain comma-separated list of
# every WG meeting that week ("CT1, CT3, CT4, CT6, RAN1, ... will be
# meeting."). Matched only when this meeting's own WG short name actually
# appears in that list, so an unrelated group's letter can't be misapplied.
_JOINT_WEEK_RE = re.compile(
    r"held\s+([A-Za-z]+\s+\d{1,2}\s*[-–]\s*\d{1,2},?\s*\d{4})\s*,?\s+in\s+([A-Za-z][^.\n]*?)\.",
    re.IGNORECASE,
)


def _split_location(raw: str):
    parts = [p.strip() for p in raw.split(",") if p.strip()]
    if not parts:
        return None, None
    location = parts[0]
    country = ", ".join(parts[1:]) if len(parts) > 1 else None
    return location, country


def _parse_table_rows(text: str, targets: set) -> Optional[MeetingInfo]:
    """Structured "TITLE  TYPE  DATES  LOCATION  COUNTRY" table, as used by
    ETSI-hosted joint CT/SA invitations — one row per co-located meeting.

    Column layout varies by invitation: a "meeting type" column (e.g. "OR")
    before the dates, and a separate country column after the location, are
    each sometimes present and sometimes not (observed: 4-part rows with no
    country column, 2-part rows with no location at all). Rather than
    assume a fixed position for dates/location/country, this scans the
    parts after the title for whichever one actually parses as a date
    range, then treats whatever comes after that as location (and country,
    if anything's left after that)."""
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        parts = [p.strip() for p in _SPLIT_RE.split(line) if p.strip()]
        if len(parts) < 2:
            continue
        row_title = _normalize_title(parts[0])
        if row_title not in targets:
            continue

        start_date = end_date = None
        date_idx = None
        for i in range(1, len(parts)):
            start_date, end_date = _parse_date_range(parts[i])
            if start_date is not None:
                date_idx = i
                break
        if date_idx is None:
            continue  # this row didn't actually look like a meeting row

        # Filter out cells that clearly aren't a real place name — one
        # invitation (RAN1#126-bis) had a "Register Here" hyperlink label
        # baked into the table where a location column would be, which
        # would otherwise get stored as if it were the venue.
        rest = [p for p in parts[date_idx + 1:] if not _JUNK_CELL_RE.search(p)]
        location = rest[0] if rest else None
        country = rest[1] if len(rest) > 1 else None
        return MeetingInfo(start_date=start_date, end_date=end_date, location=location, country=country)
    return None


def _parse_title_date_lines(text: str, targets: set) -> Optional[MeetingInfo]:
    lines = text.splitlines()
    for i, raw_line in enumerate(lines):
        stripped = raw_line.strip()
        m = _TITLE_DATE_LINE_RE.match(stripped)
        if not m:
            continue
        token, rest = m.groups()
        if _normalize_title(token) not in targets:
            continue
        start_date, end_date = _parse_date_range(rest)
        if start_date is None:
            continue

        location = country = None
        for nxt_line in lines[i + 1:]:
            nxt = nxt_line.strip()
            if not nxt:
                continue
            if _TITLE_DATE_LINE_RE.match(nxt):
                continue  # another TSG's own title/date line — keep scanning
            if not re.search(r"https?://", nxt) and "," in nxt and not re.search(r"\d{4}", nxt):
                location, country = _split_location(nxt)
            break
        return MeetingInfo(start_date=start_date, end_date=end_date, location=location, country=country)
    return None


# A host-organized (non-ETSI, non-ATIS) invitation letter, observed for
# SA4#136 (InterDigital-hosted, Montreal): a three-line header — "Invitation
# to the 3GPP <TITLE> Meeting", then the date range alone on the next line,
# then "In <LOCATION>" on the line after — rather than a table or a single
# combined title+date line.
_INVITATION_HEADER_RE = re.compile(r"Invitation\s+to\s+the\s+(?:3GPP\s*)?(.+?)\s+Meeting", re.IGNORECASE)
_IN_LOCATION_RE = re.compile(r"^in\s+(.+)$", re.IGNORECASE)


def _parse_invitation_header_block(text: str, targets: set) -> Optional[MeetingInfo]:
    lines = text.splitlines()
    for i, raw_line in enumerate(lines):
        m = _INVITATION_HEADER_RE.search(raw_line.strip())
        if not m or _normalize_title(m.group(1)) not in targets:
            continue

        start_date = end_date = None
        j = i + 1
        while j < len(lines):
            nxt = lines[j].strip()
            j += 1
            if not nxt:
                continue
            start_date, end_date = _parse_date_range(nxt)
            break
        if start_date is None:
            continue

        location = country = None
        while j < len(lines):
            nxt = lines[j].strip()
            j += 1
            if not nxt:
                continue
            loc_m = _IN_LOCATION_RE.match(nxt)
            if loc_m:
                location, country = _split_location(loc_m.group(1))
            break
        return MeetingInfo(start_date=start_date, end_date=end_date, location=location, country=country)
    return None


def _parse_joint_week_prose(text: str, expected_titles) -> Optional[MeetingInfo]:
    wg_tokens = set()
    for t in expected_titles:
        base = re.sub(r"(?i)^3GPP", "", t).split("#")[0].strip()
        if base:
            wg_tokens.add(base)
    if not wg_tokens:
        return None

    for m in _JOINT_WEEK_RE.finditer(text):
        date_text, location_text = m.groups()
        if not any(re.search(rf"\b{re.escape(tok)}\b", text) for tok in wg_tokens):
            continue
        start_date, end_date = _parse_date_range(date_text)
        if start_date is None:
            continue
        location, country = _split_location(location_text)
        return MeetingInfo(start_date=start_date, end_date=end_date, location=location, country=country)
    return None


def parse_invitation_text(text: str, expected_titles) -> Optional[MeetingInfo]:
    """expected_titles: one title string, or a list of acceptable
    variants (e.g. with/without a "3GPP" prefix, with/without a "bis"
    suffix — both observed to vary independently of the meeting-folder
    name).

    Tries four invitation letter formats, most specific first: a
    structured per-meeting table (ETSI-hosted), a per-TSG "TITLE DATE"
    line (ATIS-hosted TSG-plenary week), a host-organized three-line
    header block ("Invitation to the 3GPP <TITLE> Meeting" / date / "In
    <location>"), then a shared prose paragraph with no per-group
    breakdown at all (ATIS-hosted WG week) — the last resort since it
    can't pin dates to one specific group any more precisely than "this
    group's name is listed as meeting that week"."""
    if isinstance(expected_titles, str):
        expected_titles = [expected_titles]
    targets = {_normalize_title(t) for t in expected_titles}

    return (
        _parse_table_rows(text, targets)
        or _parse_title_date_lines(text, targets)
        or _parse_invitation_header_block(text, targets)
        or _parse_joint_week_prose(text, expected_titles)
    )


def _extract_docx_text(path: Path) -> Optional[str]:
    import docx
    from .extract import _dedupe_consecutive

    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = _dedupe_consecutive([c.text.strip() for c in row.cells if c.text.strip()])
            if cells:
                # Tab-separated, matching what LibreOffice's .doc->.txt
                # conversion naturally produces for tables — _SPLIT_RE
                # (in this same module) only recognizes tabs/double-spaces
                # as column breaks, not " | ", so joining with a pipe here
                # silently broke every .docx-format invitation's table row
                # (a plain single-space-padded pipe never split at all).
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _text_from_doc(doc_path: Path) -> Optional[str]:
    ext = doc_path.suffix.lower()
    if ext == ".docx":
        return _extract_docx_text(doc_path)
    if ext == ".pdf":
        with open(doc_path, "rb") as f:
            return _extract_pdf(f)
    return convert_doc_to_txt(doc_path)


def fetch_meeting_info(client, invitation_entries, expected_titles) -> Optional[MeetingInfo]:
    """expected_titles: passed straight through to parse_invitation_text
    — a single title string or a list of acceptable variants.

    invitation_entries: directory listing of a meeting's Invitation/
    folder (may be empty for an upcoming meeting with no invite yet).
    The invitation is sometimes a zip (CT/SA observed), sometimes a bare
    .doc/.docx/.pdf sitting directly in the folder (RAN observed) — all
    handled, resolving to a local .doc/.docx/.pdf either way before
    parsing. (RAN3#132's invitation, for instance, is a plain PDF.)"""
    candidates = [e for e in invitation_entries if not e.is_dir]
    if not candidates:
        return None

    with tempfile.TemporaryDirectory() as tmp_dir:
        tmp_dir_path = Path(tmp_dir)
        for entry in candidates:
            name_lower = entry.name.lower()
            try:
                if name_lower.endswith(".zip"):
                    zip_path = tmp_dir_path / entry.name
                    client.download(entry.url, zip_path)
                    with zipfile.ZipFile(zip_path) as z:
                        # A zip can bundle unrelated documents alongside the
                        # real invitation (observed: SA4#136's zip also had a
                        # Canada visa/immigration info letter) — trying only
                        # the first name in the archive silently parsed that
                        # instead and returned nothing. Every matching doc in
                        # the zip is tried, "invit" in the filename first
                        # since that's virtually always the real one.
                        doc_names = sorted(
                            (n for n in z.namelist() if Path(n).suffix.lower() in _INVITATION_EXTENSIONS),
                            key=lambda n: "invit" not in n.lower(),
                        )
                        for main_name in doc_names:
                            doc_path = tmp_dir_path / Path(main_name).name
                            with z.open(main_name) as src, open(doc_path, "wb") as dst:
                                shutil.copyfileobj(src, dst)
                            text = _text_from_doc(doc_path)
                            if text is None:
                                continue
                            info = parse_invitation_text(text, expected_titles)
                            if info is not None:
                                return info
                    continue
                elif name_lower.endswith(_INVITATION_EXTENSIONS):
                    doc_path = tmp_dir_path / entry.name
                    client.download(entry.url, doc_path)
                else:
                    continue

                text = _text_from_doc(doc_path)
                if text is None:
                    continue
                info = parse_invitation_text(text, expected_titles)
                if info is not None:
                    return info
            except Exception:
                log.exception("Failed to process invitation %s", entry.url)
    return None
