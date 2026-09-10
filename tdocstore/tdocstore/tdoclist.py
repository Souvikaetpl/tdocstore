"""Parse the MCC-provided TDoc_List_Meeting_<WG>#<n>.xlsx found in each meeting's Docs folder.

Column names have drifted across years; matching is case-insensitive on normalised header text.
Unknown columns are ignored; missing columns yield None. Verify against a real file for a new WG.
"""
from __future__ import annotations
import re
from datetime import datetime
from pathlib import Path
from typing import Iterator

import openpyxl

# canonical_field -> list of header aliases (normalised: lowercase, non-alnum stripped)
HEADER_ALIASES: dict[str, list[str]] = {
    "tdoc_id":        ["tdoc", "tdocnumber", "tdocid", "document", "docnumber", "number"],
    "title":          ["title"],
    "source":         ["source", "sourceorganisations", "sourceorganizations"],
    "contact":        ["contact", "contactname"],
    "document_type":  ["type", "tdoctype", "doctype"],
    "document_for":   ["for", "documentfor"],
    "agenda_item":    ["agendaitem", "agendaitemnumber", "ai", "agenda"],
    "agenda_desc":    ["agendaitemdescription", "agendaitemdesc", "agendadescription"],
    "tdoc_status":    ["tdocstatus", "status"],
    "is_revision_of": ["isrevisionof", "revisionof", "originaltdoc"],
    "revised_to":     ["revisedto", "revisedin"],
    "release":        ["release", "rel"],
    "spec":           ["spec", "specification", "specnumber"],
    "spec_version":   ["version", "specversion"],
    "related_wis":    ["relatedwis", "relatedwi", "wis", "workitem", "workitems"],
    "cr_number":      ["cr", "crnumber"],
    "cr_category":    ["crcategory", "category"],
    "uploaded_at":    ["uploaded", "uploadedon", "uploaddate", "reservationdate"],
}

TDOC_RE = re.compile(r"^[A-Z]{1,2}[0-9]?-\d{5,7}$")  # R2-2604936, RP-252912, S2-2600001, C1-260001


def _norm(h) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def _map_headers(headers: list) -> dict[str, int]:
    normed = [_norm(h) for h in headers]
    out: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        for a in aliases:
            if a in normed:
                out[field] = normed.index(a)
                break
    return out


def _cell(row, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    v = row[idx]
    if isinstance(v, datetime):
        return v.isoformat(timespec="seconds")
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def parse_tdoc_list(path: Path | str) -> Iterator[dict]:
    """Yield one dict per TDoc row. Header row is auto-detected within the first 10 rows."""
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header_map: dict[str, int] | None = None
    for i, row in enumerate(rows):
        if header_map is None:
            cand = _map_headers(list(row))
            if "tdoc_id" in cand and "title" in cand:
                header_map = cand
            elif i > 10:
                raise ValueError(f"no header row with TDoc/Title found in {path}")
            continue
        rec = {f: _cell(row, header_map.get(f)) for f in HEADER_ALIASES}
        tid = rec.get("tdoc_id")
        if not tid or not TDOC_RE.match(tid):
            continue
        # Normalise agenda item: '9.3.3.2' stays; some sheets store as float 9.3 -> keep string
        if rec.get("agenda_item") is not None:
            rec["agenda_item"] = str(rec["agenda_item"]).strip()
        yield rec
    wb.close()
